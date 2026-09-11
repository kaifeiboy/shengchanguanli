# -*- coding: utf-8 -*-
"""
inventory_export.py —— 库存尾数表格导出（openpyxl）
命令：export --rows <rows.json> --out <out.xlsx>
  rows.json: [{"kuHao":"..","code":"..","materialInfo":"..","qty":N,"createdAt":".."}, ...]
  stdout: {"ok":true,"count":N} / {"error":"..."}
"""
import os, sys, json
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except Exception as e:
    print(json.dumps({"error": "依赖不可用: %s" % e}, ensure_ascii=False)); sys.exit(1)

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADERS = ["库号", "编码", "材料信息", "当前库存", "创建时间"]


def main():
    if len(sys.argv) < 2 or sys.argv[1] != "export":
        print(json.dumps({"error": "usage: inventory_export.py export --rows <json> --out <xlsx>"}, ensure_ascii=False)); sys.exit(2)
    p = sys.argv[2:]

    def val(name, default=None):
        for i in range(0, len(p) - 1):
            if p[i] == name:
                return p[i + 1]
        return default

    rows_path = val("--rows")
    out_path = val("--out")
    if not rows_path or not out_path:
        print(json.dumps({"error": "缺少 --rows / --out"}, ensure_ascii=False)); sys.exit(2)

    with open(rows_path, "r", encoding="utf-8") as f:
        rows = json.load(f)
    if not isinstance(rows, list):
        print(json.dumps({"error": "rows 需要是数组"}, ensure_ascii=False)); sys.exit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = "库存尾数"
    widths = [18, 20, 45, 12, 22]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for r, it in enumerate(rows, start=2):
        vals = [it.get("kuHao", ""), it.get("code", ""), it.get("materialInfo", ""),
                it.get("qty", 0), it.get("createdAt", "")]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER

    try:
        wb.save(out_path)
    except Exception as e:
        print(json.dumps({"error": "保存失败: %s" % e}, ensure_ascii=False)); sys.exit(1)
    print(json.dumps({"ok": True, "count": len(rows), "path": out_path}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as ex:
        print(json.dumps({"error": "异常: %s" % ex}, ensure_ascii=False)); sys.exit(1)
