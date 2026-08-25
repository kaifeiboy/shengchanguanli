#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""修复汇总表所有有图行的行高（按图片纵横比适配，与 merge_in_cell 公式一致）。

用法: python -E _fix_row_heights.py [root 默认 E:\\生产不良履历]
走 write_book 安全流程：load → 设行高 → save .tmp → safe_replace → inject_in_cell 重建 in-cell。
图片字节原样保留（无损）。
"""
import os
import sys
import io

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import defect_history as dh
from PIL import Image


def main():
    root = sys.argv[1] if len(sys.argv) > 1 else r"E:\生产不良履历"
    book = dh.book_path(root)
    ic = dh.in_cell_images(book, dh.SHEET_NAME)
    print("in-cell 图条数:", len(ic))

    def mutate(wb):
        ws = wb[dh.SHEET_NAME]
        fixed = 0
        for d, b, f in ic:
            if not b:
                continue
            try:
                pw, ph = Image.open(io.BytesIO(b)).size
                ws.row_dimensions[d + 2].height = dh.img_disp_h_pt(pw, ph)
                fixed += 1
            except Exception:
                pass
        print("已适配行高:", fixed)
        dh.schedule_in_cell_restore(wb, ic)

    dh.write_book(root, mutate)
    print("写入完成")

    # 验证：重读行高
    from openpyxl import load_workbook
    wb2 = load_workbook(book)
    ws2 = wb2[dh.SHEET_NAME]
    for d, b, f in ic:
        h = ws2.row_dimensions[d + 2].height
        if not h:
            print("  [WARN] idx=%d 行高未设置" % d)
    print("验证完成，总数据行:", ws2.max_row - 1)
    wb2.close()


if __name__ == "__main__":
    main()
