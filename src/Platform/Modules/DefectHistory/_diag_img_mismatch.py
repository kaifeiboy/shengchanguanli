# -*- coding: utf-8 -*-
"""诊断：对比 query 返回的 hasImage/imgCount 与 xlsx 实际图片锚点行号。"""
import sys, json, subprocess
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

wb = load_workbook(r'E:\生产不良履历\不良履历汇总.xlsx'); ws = wb.active

# 实际图片行（0-based → 1-based）
img_rows = {}
for im in ws._images:
    r = im.anchor._from.row + 1
    img_rows.setdefault(r, 0)
    img_rows[r] += 1

# 数据行遍历，找 HYXC-VC01B 的所有行
print('=== xlsx 实际数据 + 图片行号 ===')
hits = []
for r in range(2, ws.max_row + 1):
    d = ws.cell(r, 1).value
    line = ws.cell(r, 2).value
    model = ws.cell(r, 3).value
    if model == 'HYXC-VC01B':
        n_img = img_rows.get(r, 0)
        hits.append((r, str(d)[:10], line, model, n_img))
        print('  row%-3d %-10s line=%-6s model=%-14s imgs=%d' % (r, str(d)[:10], line, model, n_img))

print()
print('=== /api/defecthistory/query?model=HYXC-VC01B 返回 ===')
r = subprocess.run([r'C:\Users\Administrator\.workbuddy\binaries\python\envs\default\Scripts\python.exe',
                    '-c',
                    'import urllib.request,json; r=urllib.request.urlopen("http://127.0.0.1:5000/api/defecthistory/query?model=HYXC-VC01B").read(); d=json.loads(r); print("total=",d.get("total")); print("file=",d.get("file")); [print("  row",x["row"],"date=",x["date"],"line=",x["line"],"model=",x["model"],"hasImage=",x["hasImage"],"imgCount=",x["imgCount"]) for x in d.get("rows",[])]'],
                   capture_output=True, text=True)
with open(r'E:\workaaa\shengchanguanli\temp\diag_query.txt', 'w', encoding='utf-8') as f:
    f.write(r.stdout)
print(r.stdout)
wb.close()