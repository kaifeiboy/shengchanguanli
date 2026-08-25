# -*- coding: utf-8 -*-
"""V-1 冒烟：subprocess.run 生产方式，在 defecthistory_v2_smoke 副本上验证全部命令。"""
import os, sys, io, json, shutil, subprocess, datetime
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PY = r"C:\Users\Administrator\.workbuddy\binaries\python\envs\default\Scripts\python.exe"
SCRIPT = r"E:\workaaa\shengchanguanli\src\Platform\Modules\DefectHistory\defect_history.py"
SRC = r"E:\workaaa\shengchanguanli\data\defecthistory_v2_test"
SMOKE = r"E:\workaaa\shengchanguanli\data\defecthistory_v2_smoke"

def run(*argv, cwd=None):
    env = dict(os.environ)
    env.pop("PYTHONPATH", None)
    p = subprocess.run([PY, "-E", SCRIPT] + list(argv), cwd=cwd or os.path.dirname(SCRIPT),
                       env=env, capture_output=True, text=True, encoding="utf-8", errors="replace",
                       timeout=180)
    try:
        return json.loads(p.stdout)
    except Exception:
        return {"success": False, "raw": p.stdout[:500], "err": p.stderr[:500]}

def main():
    if os.path.exists(SMOKE):
        shutil.rmtree(SMOKE, ignore_errors=True)
    os.makedirs(SMOKE)
    shutil.copy2(os.path.join(SRC, "不良履历汇总.xlsx"), os.path.join(SMOKE, "不良履历汇总.xlsx"))
    # 测试图片
    from PIL import Image as PILImage
    img_path = os.path.join(SMOKE, "_t.png")
    PILImage.new("RGB", (80, 60), (200, 60, 60)).save(img_path)
    ok_all = True

    def check(name, cond, extra=""):
        nonlocal ok_all
        print("%s %s %s" % ("PASS" if cond else "FAIL", name, extra))
        if not cond:
            ok_all = False

    # 1 meta
    r = run("meta", "--root", SMOKE)
    check("meta total=32", r.get("totalRows") == 32, "months=%s" % [m["month"] for m in r.get("months", [])])
    check("meta months 含2026-03", any(m["month"] == "2026-03" for m in r.get("months", [])))

    # 2 query
    r = run("query", "--root", SMOKE, "--model", "YCWA15NCWQ")
    check("query YCWA15NCWQ total=3", r.get("total") == 3, "total=%s" % r.get("total"))
    r = run("query", "--root", SMOKE, "--model", "YCWA15NCWQ", "--month", "2026-03")
    check("query month过滤=2", r.get("total") == 2)
    rows = r.get("rows", [])
    if rows:
        check("query 日期标准格式", rows[0]["date"].startswith("2026-03-"), "date=%s" % rows[0]["date"])
        check("query 带图标记", "hasImage" in rows[0] and "imgCount" in rows[0])

    # 3 add_single（带图）
    r = run("add_single", "--root", SMOKE, "--date", "2026.8.7", "--line", "A9-9",
            "--model", "SMOKE-ADD", "--reason", "冒烟测试", "--qty", "1台",
            "--stage", "IPQC", "--handle", "作业不良", "--images", img_path)
    check("add_single ok", r.get("ok") is True, str(r.get("error", "")))
    r = run("query", "--root", SMOKE, "--model", "SMOKE-ADD")
    check("add_single 查询命中=1", r.get("total") == 1)
    new_row = r["rows"][0]["row"] if r.get("rows") else None
    check("add_single 带图=1", r.get("rows") and r["rows"][0]["imgCount"] == 1)

    # 4 保图回归：写入后总图数（21+1=22）
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(SMOKE, "不良履历汇总.xlsx"))
    n_img = sum(len(ws._images) for ws in wb.worksheets)
    wb.close()
    check("add 后图片=22", n_img == 22, "n=%d" % n_img)

    # 5 delete_row（删 SMOKE-ADD 行）
    r = run("delete_row", "--root", SMOKE, "--row", str(new_row))
    check("delete_row ok", r.get("ok") is True)
    r = run("query", "--root", SMOKE, "--model", "SMOKE-ADD")
    check("delete 后查询=0", r.get("total") == 0)
    wb = load_workbook(os.path.join(SMOKE, "不良履历汇总.xlsx"))
    n_img2 = sum(len(ws._images) for ws in wb.worksheets)
    wb.close()
    check("delete 后图片回 21", n_img2 == 21, "n=%d" % n_img2)

    # 6 batch preview（构造含 SMT 的上传文件）
    from openpyxl import Workbook
    up = os.path.join(SMOKE, "_up.xlsx")
    wb = Workbook()
    ws = wb.active; ws.title = "组装"
    ws.append(HEADERS := ["日期", "线别", "型号", "问题原因", "数量", "发生工程", "不良图片", "处理方式"])
    ws.append(["2026.8.8", "B1-1", "BATCH-A", "批量1", "2台", "QA", "", "来料不良"])
    ws.append(["2026.8.9", "B1-2", "BATCH-B", "批量2", "3台", "IPQC", "", "作业不良"])
    ws2 = wb.create_sheet("SMT")
    ws2.append(HEADERS)
    ws2.append(["2026.8.9", "S1-1", "SMT-DROP", "应被舍弃", "1", "IPQC", "", "作业不良"])
    ws2.append(["2026.8.10", "S1-2", "SMT-DROP2", "应被舍弃", "1", "QA", "", "来料不良"])
    wb.save(up)
    r = run("batch_import", "--root", SMOKE, "--upload", up, "--preview", "1")
    check("batch preview valid=2", r.get("valid") == 2, "valid=%s" % r.get("valid"))
    check("batch preview droppedSmt=2", r.get("droppedSmt") == 2, "dropped=%s" % r.get("droppedSmt"))
    # 真实导入
    r = run("batch_import", "--root", SMOKE, "--upload", up)
    check("batch 导入=2", r.get("imported") == 2, "imported=%s err=%s" % (r.get("imported"), r.get("errors")))
    r = run("meta", "--root", SMOKE)
    check("batch 后 total=34", r.get("totalRows") == 34)
    r = run("query", "--root", SMOKE, "--model", "BATCH-A")
    check("batch 查询命中=1", r.get("total") == 1)
    wb = load_workbook(os.path.join(SMOKE, "不良履历汇总.xlsx"))
    n_img3 = sum(len(ws._images) for ws in wb.worksheets)
    wb.close()
    check("batch 后图片仍 21", n_img3 == 21)

    # 7 export full
    out_x = os.path.join(SMOKE, "_export.xlsx")
    r = run("export", "--root", SMOKE, "--out", out_x, "--full", "1")
    check("export full ok rows=34", r.get("rows") == 34, "rows=%s" % r.get("rows"))
    if os.path.exists(out_x):
        wb = load_workbook(out_x)
        ws = wb.active
        hdr = [ws.cell(1, c).value for c in range(1, 9)]
        imgs = len(ws._images)
        wb.close()
        check("export 表头完整", hdr == HEADERS, str(hdr))
        check("export 图片嵌入=21", imgs == 21, "imgs=%d" % imgs)

    # 8 template
    tpl = os.path.join(SMOKE, "_tpl.xlsx")
    r = run("template", "--root", SMOKE, "--out", tpl)
    check("template ok", r.get("ok") is True and os.path.exists(tpl))

    # 9 extract_images（第一次重建；第二次 fresh 分支）
    r = run("extract_images", "--root", SMOKE)
    f = r.get("files", [{}])[0]
    check("extract 首次重建 fresh=False", f.get("fresh") is False, "fresh=%s" % f.get("fresh"))
    s = f.get("sheets", [{}])[0]
    check("extract sheet 名=不良履历", s.get("name") == "不良履历")
    n_img_side = sum(len(x.get("thumbs", [])) for x in s.get("images", []))
    check("extract 缩略图=21", n_img_side == 21, "thumbs=%d" % n_img_side)
    # 第二次：fresh 分支必须返回 sheets
    r2 = run("extract_images", "--root", SMOKE)
    f2 = r2.get("files", [{}])[0]
    check("extract 二次 fresh=True", f2.get("fresh") is True, "fresh=%s" % f2.get("fresh"))
    check("extract fresh 分支返回 sheets", len(f2.get("sheets", [])) == 1)
    s2 = f2["sheets"][0]
    n2 = sum(len(x.get("thumbs", [])) for x in s2.get("images", []))
    check("extract fresh 缩略图=21", n2 == 21, "thumbs=%d" % n2)

    # 10 analysis
    r = run("analysis", "--root", SMOKE, "--month", "2026-03")
    check("analysis 2026-03 data>0", r.get("ok") is True and len(r.get("data", [])) > 0,
          "data=%s" % json.dumps(r.get("data"), ensure_ascii=False)[:200])
    an_out = os.path.join(SMOKE, "_analysis.xlsx")
    r = run("analysis", "--root", SMOKE, "--month", "2026-03", "--out", an_out)
    check("analysis 导出 ok", r.get("ok") is True and os.path.exists(an_out))
    if os.path.exists(an_out):
        wb = load_workbook(an_out)
        ws = wb.active
        charts = getattr(ws, "_charts", [])
        wb.close()
        check("analysis xlsx 含原生图表", len(charts) >= 1, "charts=%d" % len(charts))

    # 11 analysis_auto / status
    r = run("analysis_auto", "--root", SMOKE)
    check("analysis_auto 结构", "success" in r and r.get("success") is True, "msg=%s" % r.get("message", ""))
    r = run("analysis_status", "--root", SMOKE)
    check("analysis_status ok", "available" in r)

    print("==================== 结果:", "全部通过" if ok_all else "存在失败 ====================")
    sys.exit(0 if ok_all else 1)

if __name__ == "__main__":
    main()
