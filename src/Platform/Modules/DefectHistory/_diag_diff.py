# -*- coding: utf-8 -*-
import sys, hashlib
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

wb = load_workbook(r'E:\生产不良履历\不良履历汇总.xlsx')
ws = wb.active
img_per_xlsx_row = {}
for im in ws._images:
    img_per_xlsx_row.setdefault(im.anchor._from.row + 1, []).append(im)

# 模拟 read_rows 累加（API row 字段）
api_row_to_xlsx = {}
api_row = -1
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    if all(v is None or str(v).strip() == "" for v in vals):
        continue
    if any(isinstance(v, str) and str(v).strip().startswith("合计") for v in vals):
        continue
    api_row += 1
    api_row_to_xlsx[api_row] = r

print('总数据行数（按 read_rows 累加）:', api_row + 1)

# 模拟 sheet_images 的 d 值（= r0 - 1）
img_data_idx_to_xlsx = {}
for im in ws._images:
    r0 = im.anchor._from.row
    if r0 < 1: continue
    img_data_idx_to_xlsx[r0 - 1] = im.anchor._from.row + 1  # xlsx 1-based

print('有图数据行数（按 sheet_images）:', len(img_data_idx_to_xlsx))
print()

# 关键对比：API row=13 的实际行 vs sidecar row_13 命名的行
print('=== API row=13 ===')
print('  API row=13 指向 xlsx row:', api_row_to_xlsx.get(13))
print('  sidecar row_13 指向 xlsx row:', img_data_idx_to_xlsx.get(13))
print('  → 两者是否一致:', api_row_to_xlsx.get(13) == img_data_idx_to_xlsx.get(13))

print('=== API row=21 ===')
print('  API row=21 指向 xlsx row:', api_row_to_xlsx.get(21))
print('  sidecar row_21 指向 xlsx row:', img_data_idx_to_xlsx.get(21))
print('  → 两者是否一致:', api_row_to_xlsx.get(21) == img_data_idx_to_xlsx.get(21))

print('=== API row=30 ===')
print('  API row=30 指向 xlsx row:', api_row_to_xlsx.get(30))
print('  sidecar row_30 指向 xlsx row:', img_data_idx_to_xlsx.get(30))
print('  → 两者是否一致:', api_row_to_xlsx.get(30) == img_data_idx_to_xlsx.get(30))

print()
# 找出所有不一致的图
print('=== sidecar 命名 vs API 命名的差异 ===')
mismatch = 0
for api_r, xlsx_r_api in api_row_to_xlsx.items():
    if api_r not in img_data_idx_to_xlsx:
        continue
    if xlsx_r_api != img_data_idx_to_xlsx[api_r]:
        mismatch += 1
        if mismatch <= 10:
            print('  API row=%d → xlsx row=%d，但 sidecar row_%d 对应 xlsx row=%d' %
                  (api_r, xlsx_r_api, api_r, img_data_idx_to_xlsx[api_r]))
print('不一致图片数:', mismatch)
wb.close()