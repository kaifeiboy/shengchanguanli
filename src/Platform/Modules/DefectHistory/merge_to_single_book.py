# -*- coding: utf-8 -*-
"""
merge_to_single_book.py —— 数据源底层重构：多月份 xlsx → 单工作簿单 Sheet

固定规则（后续所有调整的基础）：
  1. 数据源 = 一张工作簿（默认 E:\\生产不良履历\\不良履历汇总.xlsx）
  2. 仅一个 Sheet「不良履历」：第1行表头8列(日期|线别|型号|问题原因|数量|发生工程|不良图片|处理方式)，
     第2行起数据，按第一列(日期)升序排列；后续所有查询基于第一列日期。
  3. SMT 工作表数据一律舍弃（含上传文件中的 SMT）。
  4. 日期自动标准化：解析多种格式 → datetime 单元格（number_format=yyyy-mm-dd），5位年份容错。
  5. 图片按行迁移到 G 列，同 cell 多图水平错开。

用法：
  python merge_to_single_book.py --src <月度xlsx目录> --out <汇总xlsx路径> [--report <json>]
输出：stdout JSON {success, keptRows, droppedSmtRows, images, dateFixes:[...], unparsed:[...], out}
"""
import os, sys, json, glob, datetime, io
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter
except Exception as e:
    print(json.dumps({"success": False, "error": "openpyxl 不可用: %s" % e}, ensure_ascii=False)); sys.exit(1)

HEADERS = ["日期", "线别", "型号", "问题原因", "数量", "发生工程", "不良图片", "处理方式"]
SHEET_NAME = "不良履历"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def parse_date(v):
    """日期标准化：返回 (datetime.date|None, 原始值)。None 表示无法解析。"""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, v
    if isinstance(v, datetime.datetime):
        return v.date(), v
    if isinstance(v, datetime.date):
        return v, v
    if isinstance(v, (int, float)):
        # Excel 日期序列号
        try:
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))), v
        except Exception:
            return None, v
    s = str(v).strip()
    if not s:
        return None, v
    # 去掉时间部分
    for sep in (" ", "T"):
        if sep in s:
            s = s.split(sep)[0]
    # 统一分隔符
    s2 = s.replace("/", ".").replace("-", ".").replace("年", ".").replace("月", ".").replace("日", ".")
    parts = [p for p in s2.split(".") if p.strip()]
    if len(parts) < 3:
        return None, v
    y, mo, d = parts[0], parts[1], parts[2]
    # 5 位年份容错：20226.3.14 → 20 + "26" = 2026
    if len(y) == 5 and y.startswith("20") and y[2] == "2":
        y = y[:2] + y[3:]
    try:
        return datetime.date(int(y), int(mo), int(d)), v
    except ValueError:
        return None, v

def sheet_images_by_row(ws):
    """返回 {data_idx: [(bdata, fmt)]}，data_idx 为数据行索引（0-based，对应 openpyxl row=3+idx）。"""
    out = {}
    for im in getattr(ws, "_images", []) or []:
        try:
            a = im.anchor
            r0 = a._from.row if a._from is not None else 0
        except Exception:
            r0 = 0
        idx = max(0, r0 - 2)  # 锚点行(0-based) → 数据行索引
        bdata = im._data()
        fmt = "PNG"
        if bdata[:8] == b"\x89PNG\r\n\x1a\n":
            fmt = "PNG"
        elif bdata[:2] == b"\xff\xd8":
            fmt = "JPEG"
        elif bdata[:2] == b"BM":
            fmt = "BMP"
        out.setdefault(idx, []).append((bdata, fmt))
    return out

def extract_sheet(root_dir, src_path):
    """读一个源文件：返回 (kept_rows, smt_dropped, images, date_fixes, unparsed)。"""
    wb = load_workbook(src_path)
    kept_rows = []      # (date_or_raw, [8列值])
    smt_dropped = 0
    images = []         # (data_idx, bdata, fmt)
    date_fixes = []
    unparsed = []
    for ws in wb.worksheets:
        if ws.title == "SMT":
            # 统计 SMT 数据行（报告用）
            for r in range(3, ws.max_row + 1):
                if any(ws.cell(r, c).value not in (None, "") for c in range(1, 9)):
                    smt_dropped += 1
            continue
        imgs = sheet_images_by_row(ws)
        for r in range(3, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, 9)]
            if not any(x is not None and str(x).strip() for x in vals):
                continue
            dval, orig = parse_date(vals[0])
            if dval is not None:
                row_rec = [dval] + vals[1:]
            else:
                row_rec = [str(vals[0]).strip() if vals[0] is not None else ""] + vals[1:]
                unparsed.append({"file": os.path.basename(src_path), "sheet": ws.title,
                                 "row": r, "raw": str(vals[0])})
            kept_rows.append(row_rec)
            idx = r - 3
            for bdata, fmt in imgs.get(idx, []):
                images.append((len(kept_rows) - 1, bdata, fmt))
            # 日期修正报告（5位年份等非标准源）
            if dval is not None and not isinstance(vals[0], (datetime.datetime, datetime.date)):
                try:
                    if isinstance(vals[0], str) and vals[0].strip() and vals[0].strip() != dval.strftime("%Y.%m.%d"):
                        date_fixes.append({"file": os.path.basename(src_path), "sheet": ws.title,
                                           "row": r, "raw": str(vals[0])[:40], "std": str(dval)})
                except Exception:
                    pass
    wb.close()
    return kept_rows, smt_dropped, images, date_fixes, unparsed

def build_book(rows, images):
    """rows=[date|raw, 8列]; images=[(final_data_idx, bdata, fmt)] → 写新工作簿。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    widths = [14, 10, 18, 40, 10, 10, 22, 14]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 22
    for i, rec in enumerate(rows):
        xr = 2 + i
        is_date = isinstance(rec[0], (datetime.date, datetime.datetime))
        ws.cell(xr, 1, rec[0] if is_date else (rec[0] if rec[0] else None))
        if is_date:
            ws.cell(xr, 1).number_format = "yyyy-mm-dd"
        for c in range(2, 9):
            v = rec[c - 1] if c - 1 < len(rec) else ""
            ws.cell(xr, c, v if v not in (None, "") else None)
        for c in range(1, 9):
            ws.cell(xr, c).border = BORDER
    # 图片：按数据行分组，同 cell 多图水平错开
    by_row = {}
    for idx, bdata, fmt in images:
        by_row.setdefault(idx, []).append(bdata)
    for idx, blist in sorted(by_row.items()):
        xr = 2 + idx
        for k, bdata in enumerate(blist):
            im = XlImage(io.BytesIO(bdata))
            # 限制最大宽度，保持比例
            if im.width > 200:
                ratio = 200.0 / im.width
                im.width = int(im.width * ratio)
                im.height = int(im.height * ratio)
            col_off = k * 40 * 9525  # 每张水平偏移 40px
            # 数据行 idx（0-based）在 1-based row (2+idx)；"G{n}" 的 _from.row = n-1 → 传 G{2+idx} 恰好锚到数据行
            anchor = "G%d" % (2 + idx)
            try:
                ws.add_image(im, anchor)
                if k > 0:
                    im.anchor._from.colOff = col_off
            except Exception:
                pass
    return wb

def main(argv):
    args = {}
    i = 0
    while i < len(argv):
        a = argv[i]
        if a.startswith("--") and "=" in a:
            k, v = a.split("=", 1)
            args[k[2:].replace("-", "_")] = v
            i += 1
        elif a.startswith("--"):
            args[a[2:].replace("-", "_")] = argv[i + 1] if i + 1 < len(argv) else "1"
            i += 2
        else:
            i += 1
    src_dir = args.get("src")
    out_path = args.get("out")
    if not src_dir or not out_path:
        print(json.dumps({"success": False, "error": "缺少 --src 或 --out"}, ensure_ascii=False)); return 1
    files = sorted(glob.glob(os.path.join(src_dir, "*.xlsx")))
    all_rows, all_imgs, date_fixes, unparsed = [], [], [], []
    smt_total = 0
    per_file = []
    for f in files:
        rows, smt, imgs, fixes, unpars = extract_sheet(src_dir, f)
        offset = len(all_rows)
        all_rows.extend(rows)
        all_imgs.extend((offset + idx, bdata, fmt) for idx, bdata, fmt in imgs)
        smt_total += smt
        date_fixes.extend(fixes)
        unparsed.extend(unpars)
        per_file.append({"file": os.path.basename(f), "kept": len(rows),
                         "smtDropped": smt, "images": len(imgs)})
    # 日期排序：可解析日期在前升序，未解析的原始文本在后
    dated = [(r, i) for i, r in enumerate(all_rows) if isinstance(r[0], (datetime.date, datetime.datetime))]
    raw = [(r, i) for i, r in enumerate(all_rows) if not isinstance(r[0], (datetime.date, datetime.datetime))]
    dated.sort(key=lambda t: t[0][0])
    ordered = [r for r, _ in dated] + [r for r, _ in raw]
    # 重映射图片索引
    idx_map = {}
    for new_i, (r, old_i) in enumerate(dated + raw):
        idx_map[old_i] = new_i
    remapped = [(idx_map[old_i], bdata, fmt) for old_i, bdata, fmt in all_imgs]
    wb = build_book(ordered, remapped)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    wb.close()
    res = {"success": True, "out": out_path, "keptRows": len(ordered),
           "droppedSmtRows": smt_total, "images": len(remapped),
           "unparsed": unparsed, "dateFixes": date_fixes, "perFile": per_file}
    if args.get("report"):
        with open(args["report"], "w", encoding="utf-8") as fh:
            json.dump(res, fh, ensure_ascii=False, indent=1, default=str)
    print(json.dumps(res, ensure_ascii=False, default=str))
    return 0

if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
