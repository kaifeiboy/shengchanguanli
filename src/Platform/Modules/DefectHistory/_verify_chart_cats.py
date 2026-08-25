# -*- coding: utf-8 -*-
"""验证分析导出图表：类别引用与系列值行级对齐、与数据表合计交叉核对。"""
import sys, re
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else r'E:\workaaa\shengchanguanli\temp\an_fixed.xlsx'
wb = load_workbook(path)
ws = wb.active
ch = getattr(ws, '_charts', [None])[0]
print('grouping:', ch.grouping, '| series 数:', len(ch.series))

REF_RE = re.compile(r"^'?([^']+)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$")

def expand(ref):
    if not ref:
        return []
    m = REF_RE.match(ref)
    if not m:
        print('  [ref parse fail]', ref)
        return []
    sheet, c1, r1, c2, r2 = m.group(1), m.group(2), int(m.group(3)), m.group(4), int(m.group(5))
    wsx = wb[sheet] if sheet in wb.sheetnames else ws
    col1 = ord(c1) - 64
    col2 = ord(c2) - 64
    out = []
    for r in range(r1, r2 + 1):
        row = tuple(wsx.cell(r, c).value for c in range(col1, col2 + 1))
        out.append(row[0] if col1 == col2 else row)
    return out

cats = ch.series[0].cat
catref = cats.numRef.f if cats and cats.numRef else (cats.strRef.f if cats and cats.strRef else None)
print('类别引用:', catref)
cats_vals = expand(catref)
print('类别展开:', cats_vals)

print('--- 系列→值（X 标签 | 各型号列值）---')
for si, s in enumerate(ch.series):
    vref = s.val.numRef.f if s.val and s.val.numRef else None
    vals = expand(vref)
    tx = s.tx.strRef.f.split('!')[-1] if s.tx and s.tx.strRef else '?'
    print('  %-12s %s → %s' % (tx, vref, vals))

print('--- 数据表合计（正确基准）---')
tb = {}
for r in range(3, ws.max_row + 1):
    line, model, tot = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 5).value
    if line and model and tot is not None:
        tb.setdefault(line, {})[model] = tot
for line in cats_vals or []:
    print('  %-6s %s' % (line, {k: v for k, v in sorted(tb.get(line, {}).items())}))

# 交叉校验：类别数与系列值行数一致
if cats_vals:
    n = len(cats_vals)
    print('--- 对齐校验 ---')
    ok = True
    for si, s in enumerate(ch.series):
        vref = s.val.numRef.f if s.val and s.val.numRef else None
        m = REF_RE.match(vref or '')
        if m:
            r1, r2 = int(m.group(3)), int(m.group(5))
            if (r2 - r1 + 1) != n:
                ok = False
                print('  series%d 行数(%d) != 类别数(%d)' % (si, r2 - r1 + 1, n))
    # 值行内容逐行核对：类别[i] 拉线的该型号合计 == 系列值[i]
    for si, s in enumerate(ch.series):
        vref = s.val.numRef.f if s.val and s.val.numRef else None
        vals = expand(vref)
        txc = s.tx.strRef.f if s.tx and s.tx.strRef else ''
        model = ''
        if '!' in txc:
            sh, cell = txc.split('!', 1)
            cm = re.match(r'\$?([A-Z]+)\$?(\d+)$', cell)
            if cm:
                wsx = wb[sh] if sh in wb.sheetnames else ws
                model = wsx.cell(int(cm.group(2)), ord(cm.group(1)) - 64).value or ''
        for i, line in enumerate(cats_vals):
            expect = tb.get(line, {}).get(model)
            got = vals[i] if i < len(vals) else None
            if (expect or 0) != (got or 0):
                ok = False
                print('  MISMATCH 类别[%d]=%s 型号=%s 应=%s 实=%s' % (i, line, model, expect, got))
    print('对齐校验:', 'OK — 每个类别(拉线)标签与柱体(该拉线各型号单数)一一对应' if ok else 'FAILED')
wb.close()
