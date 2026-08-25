# -*- coding: utf-8 -*-
"""把源 xlsx 的图片保存出来，对比看源数据是否本身就有错位。"""
import sys, os, hashlib
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

out_dir = r'E:\workaaa\shengchanguanli\temp\src_imgs'
os.makedirs(out_dir, exist_ok=True)
wb = load_workbook(r'E:\生产不良履历\不良履历汇总.xlsx')
ws = wb.active

# 保存所有图片按 (sheet, xlsx_row)
img_per_row = {}
for im in ws._images:
    data_bytes = im._data()
    row1b = im.anchor._from.row + 1
    img_per_row.setdefault(row1b, []).append((data_bytes, im.format or 'png'))

# 输出每个有图行的内容 + 保存图
print('=== 源 xlsx 所有有图行 ===')
for r in sorted(img_per_row.keys()):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    md5s = [hashlib.md5(d).hexdigest()[:10] for d, _ in img_per_row[r]]
    print('  r%-3d | %-12s | %-8s | %-12s | 图=%d %s' %
          (r, str(vals[0])[:10], str(vals[1] or ''), str(vals[2] or '')[:12], len(img_per_row[r]), md5s))
    for k, (bdata, fmt) in enumerate(img_per_row[r]):
        ext = 'jpg' if fmt.lower() == 'jpeg' else fmt.lower()
        out = os.path.join(out_dir, 'src_r%d_%d.%s' % (r, k, ext))
        with open(out, 'wb') as fh:
            fh.write(bdata)
wb.close()
print('saved to', out_dir)