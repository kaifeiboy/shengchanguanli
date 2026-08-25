# -*- coding: utf-8 -*-
"""对比修正前后导出图表的 X 轴类别引用与展开值。"""
import sys, re
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

REF_RE = re.compile(r"^'?([^']+)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$")

def dump(path, tag):
    wb = load_workbook(path)
    ws = wb.active
    ch = getattr(ws, '_charts', [None])[0]
    cats = ch.series[0].cat
    ref = cats.numRef.f if cats and cats.numRef else (cats.strRef.f if cats and cats.strRef else '')
    m = REF_RE.match(ref)
    print('%s %s' % (tag, path))
    print('  类别引用:', ref)
    if m:
        wsx = wb[m.group(1)] if m.group(1) in wb.sheetnames else ws
        vals = [wsx.cell(r, ord(m.group(2)) - 64).value for r in range(int(m.group(3)), int(m.group(5)) + 1)]
        print('  X 轴标签展开:', vals)
    # 期望标签（数据表 A3.. 唯一拉线，顺序=lines）
    lines = []
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and v not in lines:
            lines.append(v)
    print('  期望标签(数据表拉线):', lines)
    print('  一致:', (vals == lines) if m else False)
    wb.close()

dump(r'C:\Users\Administrator\Desktop\不良分析_2026-03.xlsx', '[修正前]')
dump(r'E:\workaaa\shengchanguanli\temp\an_fixed.xlsx', '[修正后]')
