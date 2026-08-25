# -*- coding: utf-8 -*-
import sys, hashlib, os
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from openpyxl import load_workbook

wb = load_workbook(r'E:\生产不良履历\不良履历汇总.xlsx')
ws = wb.active
img_per_xlsx_row = {}
for im in ws._images:
    img_per_xlsx_row.setdefault(im.anchor._from.row + 1, []).append(im)

# 按 read_rows 逻辑遍历（从 row 2 起、跳过全空行、跳过"合计"行），累加 data_idx
data_idx = -1
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    if all(v is None or str(v).strip() == "" for v in vals):
        continue
    if any(isinstance(v, str) and str(v).strip().startswith("合计") for v in vals):
        continue
    data_idx += 1
    model = vals[2]
    if model == 'HYXC-VC01B':
        n_img = len(img_per_xlsx_row.get(r, []))
        print('API row=%-3d → xlsx row=%-3d imgs=%d | %s | %s' %
              (data_idx, r, n_img, str(vals[0])[:10], vals[1]))
wb.close()

print()
print('--- MD5 对比 ---')
# 重新打开（关闭上面）
wb = load_workbook(r'E:\生产不良履历\不良履历汇总.xlsx')
ws = wb.active
data_idx = -1
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    if all(v is None or str(v).strip() == "" for v in vals):
        continue
    if any(isinstance(v, str) and str(v).strip().startswith("合计") for v in vals):
        continue
    data_idx += 1
    if vals[2] != 'HYXC-VC01B' or not img_per_xlsx_row.get(r):
        continue
    im = img_per_xlsx_row[r][0]
    xlsx_md5 = hashlib.md5(im._data()).hexdigest()[:12]
    curl_path = r'E:\workaaa\shengchanguanli\temp\img_row%d.png' % data_idx
    if os.path.exists(curl_path):
        curl_md5 = hashlib.md5(open(curl_path, 'rb').read()).hexdigest()[:12]
        ok = '✓' if xlsx_md5 == curl_md5 else '✗ 错位'
        print('  API row=%d xlsx row=%d | xlsx_md5=%s curl_md5=%s %s' % (data_idx, r, xlsx_md5, curl_md5, ok))
wb.close()