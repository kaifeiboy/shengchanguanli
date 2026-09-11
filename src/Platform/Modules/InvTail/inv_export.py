# -*- coding: utf-8 -*-
"""
库存尾数模块 · 出入库明细流水导出（xlsx）

用法：
  python inv_export.py --data <payload.json> --out <目标.xlsx>

payload.json 结构（由 C# 服务层生成）：
{
  "title": "库存尾数 · 出入库明细流水",
  "generatedAt": "2026-09-02 14:00:00",
  "filter": {"keyword": "", "handler": "", "direction": "", "dateFrom": "", "dateTo": ""},
  "rows": [{"binNo","code","materialInfo","direction","qty","balance","occurredAt","handler","note"}, ...]
}

一行一笔流水，用于追溯核账：可看到每笔出入库后的滚动结存。
stdout 仅输出一段 JSON：{"success": true/false, ...}
"""
import argparse
import io
import json
import sys

_REAL_STDOUT = sys.stdout
sys.stdout = io.StringIO()

from openpyxl import Workbook                                              # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side     # noqa: E402
from openpyxl.utils import get_column_letter                               # noqa: E402

HEADERS = ['序号', '库号', '编码', '材料信息', '方向', '数量', '结存', '时间', '经手人', '备注']
WIDTHS = [6, 16, 16, 34, 8, 10, 10, 20, 12, 22]

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='2563EB')
IN_FILL = PatternFill('solid', fgColor='E8F1FE')     # 入库：淡蓝
OUT_FILL = PatternFill('solid', fgColor='FDECEC')    # 出库：淡红
SUM_FILL = PatternFill('solid', fgColor='F5F7FA')


def fmt_num(v):
    """数量显示：整数不带小数点，小数保留有效位。"""
    try:
        f = float(v)
    except Exception:
        return v
    if abs(f - round(f)) < 1e-9:
        return int(round(f))
    return round(f, 4)


def filter_text(flt):
    if not isinstance(flt, dict):
        return '筛选条件：全部'
    parts = []
    if flt.get('keyword'):
        parts.append('关键词=%s' % flt['keyword'])
    if flt.get('handler'):
        parts.append('经手人=%s' % flt['handler'])
    d = (flt.get('direction') or '').lower()
    if d == 'in':
        parts.append('方向=入库')
    elif d == 'out':
        parts.append('方向=出库')
    if flt.get('dateFrom'):
        parts.append('起始=%s' % flt['dateFrom'])
    if flt.get('dateTo'):
        parts.append('截止=%s' % flt['dateTo'])
    return '筛选条件：' + ('，'.join(parts) if parts else '全部')


def build(payload, out_path):
    rows = payload.get('rows') or []
    title = payload.get('title') or '库存尾数 · 出入库明细流水'
    generated = payload.get('generatedAt') or ''

    wb = Workbook()
    ws = wb.active
    ws.title = '出入库明细'
    ncol = len(HEADERS)
    last_col = get_column_letter(ncol)

    # ---- 标题 ----
    ws.merge_cells('A1:%s1' % last_col)
    c = ws['A1']
    c.value = title
    c.font = Font(name='微软雅黑', size=15, bold=True, color='1F2937')
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # ---- 副标题：筛选条件 + 导出时间 ----
    ws.merge_cells('A2:%s2' % last_col)
    c = ws['A2']
    c.value = '%s        导出时间：%s        共 %d 笔' % (filter_text(payload.get('filter')), generated, len(rows))
    c.font = Font(name='微软雅黑', size=9, color='6B7280')
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 20

    # ---- 表头 ----
    head_row = 3
    for i, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=head_row, column=i, value=h)
        cell.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    ws.row_dimensions[head_row].height = 24

    # ---- 数据 ----
    total_in = 0.0
    total_out = 0.0
    r = head_row + 1
    for idx, item in enumerate(rows, start=1):
        direction = (item.get('direction') or '').strip()
        qty = item.get('qty') or 0
        try:
            q = float(qty)
        except Exception:
            q = 0.0
        if direction == '出库':
            total_out += q
        else:
            total_in += q

        values = [
            idx,
            item.get('binNo') or '',
            item.get('code') or '',
            item.get('materialInfo') or '',
            direction,
            fmt_num(qty),
            fmt_num(item.get('balance') or 0),
            item.get('occurredAt') or '',
            item.get('handler') or '',
            item.get('note') or ''
        ]
        for ci, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.border = BORDER
            cell.font = Font(name='微软雅黑', size=10, color='1F2937')
            if ci in (1, 5, 6, 7):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif ci == 4:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
        # 方向着色：入库淡蓝、出库淡红，便于肉眼扫读
        dcell = ws.cell(row=r, column=5)
        if direction == '出库':
            dcell.fill = OUT_FILL
            dcell.font = Font(name='微软雅黑', size=10, bold=True, color='C0392B')
        else:
            dcell.fill = IN_FILL
            dcell.font = Font(name='微软雅黑', size=10, bold=True, color='1D4ED8')
        r += 1

    # ---- 合计 ----
    ws.cell(row=r, column=1, value='合计')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    tc = ws.cell(row=r, column=1)
    tc.font = Font(name='微软雅黑', size=10, bold=True)
    tc.alignment = Alignment(horizontal='center', vertical='center')
    summary = '入库 %s ／ 出库 %s ／ 净增 %s' % (
        fmt_num(total_in), fmt_num(total_out), fmt_num(total_in - total_out))
    ws.cell(row=r, column=5, value=summary)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=len(HEADERS))
    sc = ws.cell(row=r, column=5)
    sc.font = Font(name='微软雅黑', size=10, bold=True, color='1F2937')
    sc.alignment = Alignment(horizontal='left', vertical='center')
    for ci in range(1, len(HEADERS) + 1):
        ws.cell(row=r, column=ci).border = BORDER
        ws.cell(row=r, column=ci).fill = SUM_FILL
    ws.row_dimensions[r].height = 22

    # ---- 列宽 / 冻结 / 筛选 ----
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A%d' % (head_row + 1)
    if rows:
        ws.auto_filter.ref = 'A%d:%s%d' % (head_row, last_col, head_row + len(rows))

    wb.save(out_path)
    return {'success': True, 'out': out_path, 'rows': len(rows),
            'totalIn': fmt_num(total_in), 'totalOut': fmt_num(total_out)}


def main():
    ap = argparse.ArgumentParser(description='库存尾数出入库明细导出')
    ap.add_argument('--data', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    with open(args.data, 'r', encoding='utf-8') as f:
        payload = json.load(f)
    return build(payload, args.out)


if __name__ == '__main__':
    try:
        result = main()
    except Exception as ex:
        result = {'success': False, 'error': str(ex)}
    finally:
        sys.stdout = _REAL_STDOUT
    print(json.dumps(result, ensure_ascii=False))
